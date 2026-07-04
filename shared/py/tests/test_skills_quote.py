"""skills/drawing_to_quotation 机械脚本测试（纯离线：合成 PDF/模板，不碰网络）。

覆盖：单元（骨架字段/款号展开/数量解析）、集成（scaffold→fill_quote 闭环）、
边界异常（require_visual 闸门/空 JSON）、端到端渲染（有 soffice 才跑）。
"""
from __future__ import annotations

import functools
import json
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

import fitz
import pytest
from openpyxl import load_workbook, Workbook

from skills.drawing_to_quotation import (
    QuoteAuditError,
    build_scaffold,
    fill_quote_build,
    render_xlsx,
)


@pytest.fixture
def drawing_pdf(tmp_path: Path) -> str:
    """两页合成图纸：第1页封面；第2页矢量视图 + 注记（含缩写款号与分品番数量）。"""
    doc = fitz.open()
    cover = doc.new_page(width=842, height=595)
    cover.insert_text((100, 100), "COVER PAGE", fontsize=24)

    page = doc.new_page(width=842, height=595)
    # 一个 400x200pt 的“产品正面图” + 底部外形尺寸线（标注 1200 → 比例尺 3mm/pt）
    page.draw_rect(fitz.Rect(100, 100, 500, 300))
    page.draw_line(fitz.Point(100, 320), fitz.Point(500, 320))
    page.insert_text((290, 315), "1200", fontsize=8)
    page.draw_line(fitz.Point(520, 100), fitz.Point(520, 300))
    page.insert_text((525, 200), "600", fontsize=8)
    note = (
        "案件名：测试项目\n产品名：テーブル-01\n款号：F01，01A\n"
        "数量：F01：1pcs。F01A：2pcs。共3pcs\nW1200 D850 H725\n材质：メラミン化粧板"
    )
    page.insert_text((550, 400), note, fontsize=7, fontname="china-s")  # 默认 helv 丢 CJK 字形
    out = tmp_path / "drawing.pdf"
    doc.save(out)
    return str(out)


@pytest.fixture
def template_xlsx(tmp_path: Path) -> str:
    wb = Workbook()
    ws = wb.active
    ws["C11"] = ""
    out = tmp_path / "template.xlsx"
    wb.save(out)
    return str(out)


# ---------- 单元：build_scaffold ----------

def test_build_scaffold_expands_codes_and_leaves_dims_null(drawing_pdf, tmp_path):
    out_json = tmp_path / "products.json"
    payload = build_scaffold(drawing_pdf, str(out_json), skip_pages=[1], project="测试项目")

    codes = [p["row_code"] for p in payload["products"]]
    assert codes == ["F01", "F01A"], "「F07，07A」式缩写款号应展开成两条记录"
    for p in payload["products"]:
        assert p["W"] is None and p["D"] is None and p["H"] is None, "骨架里 W/D/H 必须留空"
        assert p["dim_source"] == "PENDING"
        assert p["page"] == 2
        assert Path(p["page_image"]).exists(), "整页图应已渲染"
        assert p["text_dims"] == {"W": 1200, "D": 850, "H": 725}
    hints = {p["row_code"]: p["qty_hint"] for p in payload["products"]}
    assert hints == {"F01": 1, "F01A": 2}, "分品番数量应各归各解析"
    assert json.loads(out_json.read_text(encoding="utf-8"))["project"] == "测试项目"


def test_build_scaffold_skip_pages_and_placeholder(drawing_pdf, tmp_path):
    payload = build_scaffold(drawing_pdf, str(tmp_path / "p.json"), skip_pages=[2])
    # 只剩封面页：无品番 → 建空品番占位记录（无文本层页也要建行的铁律）
    assert [p["row_code"] for p in payload["products"]] == [""]
    assert payload["products"][0]["page"] == 1


# ---------- 集成：scaffold → fill_quote 闭环 ----------

def _confirm_all(products):
    for p in products:
        p.update(W=1200, D=850, H=725, qty=p["qty_hint"], dim_source="visual",
                 dim_evidence="正面図下端の外形寸法線 W1200", name_jp="テーブル-01",
                 name_cn="餐桌-01", mat_jp=["メラミン化粧板"], mat_cn=["三聚氰胺装饰板"])


def test_fill_quote_roundtrip(drawing_pdf, template_xlsx, tmp_path):
    out_json = tmp_path / "products.json"
    payload = build_scaffold(drawing_pdf, str(out_json), skip_pages=[1], project="测试项目")
    _confirm_all(payload["products"])
    out_json.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

    out_xlsx = tmp_path / "out.xlsx"
    summary = fill_quote_build(drawing_pdf, template_xlsx, str(out_json), str(out_xlsx),
                               "", 18, 30, False)
    assert summary["rows"] == 2 and summary["unconfirmed"] == [] and summary["missing"] == []
    ws = load_workbook(out_xlsx).active
    assert ws["C11"].value == "测试项目"
    assert ws["B18"].value == "F01" and ws["B19"].value == "F01A", "一个品番一行，绝不合并"
    assert ws["I18"].value == 1 and ws["I19"].value == 2, "数量各归各"
    assert "防火板" in ws["E18"].value and "三聚氰胺" not in ws["E18"].value, "术语订正应生效"


def test_fill_quote_flags_unconfirmed_rows(drawing_pdf, template_xlsx, tmp_path):
    out_json = tmp_path / "products.json"
    build_scaffold(drawing_pdf, str(out_json), skip_pages=[1])
    out_xlsx = tmp_path / "draft.xlsx"
    summary = fill_quote_build(drawing_pdf, template_xlsx, str(out_json), str(out_xlsx),
                               "", 18, 30, False)
    assert set(summary["unconfirmed"]) == {"F01", "F01A"}
    ws = load_workbook(out_xlsx).active
    # 新方案：品番不再加 ⚠ 前缀；缺尺寸/数量的行改为备注写极简标记 + 不确定单格淡黄。
    assert not (ws["B18"].value or "").startswith("⚠"), "品番不应再加 ⚠ 前缀"
    assert "不确定" in (ws["N18"].value or ""), "缺尺寸/数量的行备注应写「…不确定」"


# ---------- 边界 / 异常 ----------

def test_require_visual_gate_raises(drawing_pdf, template_xlsx, tmp_path):
    out_json = tmp_path / "products.json"
    build_scaffold(drawing_pdf, str(out_json), skip_pages=[1])  # 全 PENDING
    with pytest.raises(QuoteAuditError):
        fill_quote_build(drawing_pdf, template_xlsx, str(out_json),
                         str(tmp_path / "final.xlsx"), "", 18, 30, True)


def test_empty_products_json_raises(drawing_pdf, template_xlsx, tmp_path):
    out_json = tmp_path / "empty.json"
    out_json.write_text('{"project":"","products":[]}', encoding="utf-8")
    with pytest.raises(QuoteAuditError):
        fill_quote_build(drawing_pdf, template_xlsx, str(out_json),
                         str(tmp_path / "o.xlsx"), "", 18, 30, False)


# ---------- 端到端渲染（需 LibreOffice，本机没有则跳过；Docker 里会跑）----------


@functools.lru_cache(maxsize=1)
def _soffice_works() -> bool:
    """光有 soffice 可执行文件还不够——某些容器里它装了却转不了任何文件
    （报 "source file could not be loaded"，且退出码仍是 0）。真跑一次极小转换来判活，
    转不出 PDF 就跳过端到端渲染测试。"""
    if shutil.which("soffice") is None:
        return False
    try:
        with tempfile.TemporaryDirectory() as d:
            src = os.path.join(d, "probe.csv")
            Path(src).write_text("a,b\n1,2\n", encoding="utf-8")
            subprocess.run(
                ["soffice", "--headless", "--convert-to", "pdf", "--outdir", d, src],
                check=True, timeout=60,
                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
            )
            return os.path.exists(os.path.join(d, "probe.pdf"))
    except Exception:
        return False


@pytest.mark.skipif(not _soffice_works(), reason="LibreOffice(soffice) 不可用或无法转换文件")
def test_render_xlsx_produces_pngs(drawing_pdf, template_xlsx, tmp_path):
    out_json = tmp_path / "products.json"
    payload = build_scaffold(drawing_pdf, str(out_json), skip_pages=[1])
    _confirm_all(payload["products"])
    out_json.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    out_xlsx = tmp_path / "out.xlsx"
    fill_quote_build(drawing_pdf, template_xlsx, str(out_json), str(out_xlsx), "", 18, 30, False)
    pngs = render_xlsx(str(out_xlsx), str(tmp_path / "check"))
    assert pngs and all(Path(p).exists() for p in pngs)
