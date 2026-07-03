"""任务1「报价表行数自适应」的六维度测试。

被测：fill_quote.build（= 包 fill_quote_build）在写完 N 行数据后，把模板里的
「页脚块」（合計/船運/契约/備考）整块搬到「数据末行 + 1 空白行」之后，并修正
合計 SUM 范围末行、页脚内部公式、顶部对页脚的引用、样式/行高/合并单元格。

现有 test_skills_quote 用的合成模板没有页脚，压不到这段逻辑；这里专门造一个
**带页脚**的模板来验。

维度：单元 / 集成 / 回归 / 边界·异常 / 端到端（soffice 重算） / 黄金快照。
"""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

import fitz
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from skills.drawing_to_quotation import fill_quote_build

START_ROW = 18
LAST_ROW = 25            # 数据预留段末行 = 合計行的上一行
FS = LAST_ROW + 1        # 合計行 = 26
# 页脚：26 合計 / 27 船運 / 28 契约 / 29 備考
FOOTER_LABELS = {0: "合計", 1: "船運費", 2: "契約金額", 3: "備考：数量以実測を優先"}


# ---------------- 合成资源 ----------------
@pytest.fixture
def blank_pdf(tmp_path) -> str:
    p = tmp_path / "blank.pdf"
    doc = fitz.open()
    doc.new_page(width=595, height=842)
    doc.save(str(p))
    doc.close()
    return str(p)


def _make_footer_template(path: str) -> None:
    """造一个带页脚块的报价模板：
    - 顶部 C9 = =K{合計行}（测顶部引用平移）
    - 数据预留 18..25 空
    - 页脚 26 合計 / 27 船運 / 28 契约 / 29 備考，含公式/样式/行高/合并
    """
    wb = Workbook()
    ws = wb.active
    bold = Font(name="微软雅黑", size=11, bold=True)
    fill = PatternFill("solid", fgColor="FFDDDDDD")

    ws["C9"] = f"=K{FS}"                                   # 顶部引用合計
    # 合計行（rel 0）
    ws[f"A{FS}"] = FOOTER_LABELS[0]
    ws[f"K{FS}"] = f"=SUM(K{START_ROW}:K{LAST_ROW})"      # ← SUM 从数据起始行到预留末行
    ws[f"K{FS}"].font = bold
    ws[f"K{FS}"].fill = fill
    ws[f"K{FS}"].number_format = "0.00"
    ws.row_dimensions[FS].height = 22
    ws.merge_cells(start_row=FS, start_column=1, end_row=FS, end_column=10)  # A..J 合并作标签
    # 船運（rel 1）：内部引用合計行
    ws[f"A{FS+1}"] = FOOTER_LABELS[1]
    ws[f"K{FS+1}"] = f"=K{FS}*0.05"
    ws.row_dimensions[FS + 1].height = 18
    # 契约（rel 2）：合計 + 船運
    ws[f"A{FS+2}"] = FOOTER_LABELS[2]
    ws[f"K{FS+2}"] = f"=K{FS}+K{FS+1}"
    ws.row_dimensions[FS + 2].height = 18
    # 備考（rel 3）：纯文本 + 合并
    ws[f"A{FS+3}"] = FOOTER_LABELS[3]
    ws.row_dimensions[FS + 3].height = 40
    ws.merge_cells(start_row=FS + 3, start_column=1, end_row=FS + 3, end_column=14)
    wb.save(path)


@pytest.fixture
def footer_template(tmp_path) -> str:
    p = tmp_path / "tmpl_footer.xlsx"
    _make_footer_template(str(p))
    return str(p)


def _mk_products(n: int) -> list[dict]:
    """n 条已视觉确认、三维齐全的产品（page=None → 不裁图，便于计数）。"""
    out = []
    for i in range(n):
        out.append({
            "row_code": f"F{i + 1:02d}", "page": None,
            "W": 1000, "D": 500, "H": 700, "qty": 2,
            "name_jp": "テスト", "name_cn": "测试",
            "mat_jp": ["メラミン"], "mat_cn": ["防火板"],
            "note_jp": "", "note_cn": "",
            "dim_source": "visual", "dim_evidence": "ok",
        })
    return out


# 每条 K = F*G*H*I/1e9*1.1 = 1000*500*700*2/1e9*1.1
K_PER = 1000 * 500 * 700 * 2 / 1_000_000_000 * 1.1  # = 0.77


def _run(template: str, products: list[dict], pdf: str, tmp_path,
         last_row: int = LAST_ROW, require_visual: bool = False):
    """跑 fill_quote_build，返回 (out_path, 重新载入的 ws, summary)。"""
    jp = tmp_path / f"products_{len(products)}.json"
    jp.write_text(json.dumps({"project": "", "products": products}, ensure_ascii=False),
                  encoding="utf-8")
    out = tmp_path / f"out_{len(products)}.xlsx"
    summary = fill_quote_build(pdf, template, str(jp), str(out), "",
                               START_ROW, last_row, require_visual)
    ws = load_workbook(str(out)).active
    return str(out), ws, summary


def _layout(n: int, last_row: int = LAST_ROW):
    data_end = START_ROW + n - 1
    blank_row = data_end + 1
    new_fs = blank_row + 1
    return data_end, blank_row, new_fs


# ==================== 单元 ====================
def test_unit_footer_relocates_after_data_and_blank(footer_template, blank_pdf, tmp_path):
    """N 少于预留段：数据 N 行 → 1 空白行 → 合計，旧页脚位置被清空。"""
    n = 3
    _, ws, _ = _run(footer_template, _mk_products(n), blank_pdf, tmp_path)
    data_end, blank_row, new_fs = _layout(n)
    # 数据末行有内容，空白行真空白
    assert ws[f"B{data_end}"].value  # 品番在
    assert all(ws.cell(blank_row, c).value in (None, "") for c in range(1, 15)), "空白行应真空白"
    # 合計搬到 new_fs
    assert ws[f"A{new_fs}"].value == FOOTER_LABELS[0]
    assert ws[f"K{new_fs}"].value == f"=SUM(K{START_ROW}:K{data_end})"
    # 旧页脚位置（26..29）已清空
    for rr in range(FS, FS + 4):
        assert ws[f"A{rr}"].value in (None, ""), f"旧页脚 A{rr} 应被清空"
        assert ws[f"K{rr}"].value in (None, ""), f"旧页脚 K{rr} 应被清空"


# ==================== 集成 ====================
def test_integration_formulas_styles_merges_preserved(footer_template, blank_pdf, tmp_path):
    n = 5
    _, ws, summary = _run(footer_template, _mk_products(n), blank_pdf, tmp_path)
    data_end, blank_row, new_fs = _layout(n)
    # 合計 SUM 末行 = 真实数据末行
    assert ws[f"K{new_fs}"].value == f"=SUM(K{START_ROW}:K{data_end})"
    # 船運 / 契约 内部引用按位移平移（Translator）
    assert ws[f"K{new_fs+1}"].value == f"=K{new_fs}*0.05"
    assert ws[f"K{new_fs+2}"].value == f"=K{new_fs}+K{new_fs+1}"
    # 顶部引用平移
    assert ws["C9"].value == f"=K{new_fs}"
    # 備考文本在新位置
    assert ws[f"A{new_fs+3}"].value == FOOTER_LABELS[3]
    # 样式：合計加粗 + 灰底 + 数字格式 + 行高
    assert ws[f"K{new_fs}"].font.bold is True
    assert (ws[f"K{new_fs}"].fill.fgColor.rgb or "").endswith("DDDDDD")
    assert ws[f"K{new_fs}"].number_format == "0.00"
    assert ws.row_dimensions[new_fs].height == 22
    assert ws.row_dimensions[new_fs + 3].height == 40
    # 合并单元格在新位置重建（合計标签 A..J、備考 A..N）
    merges = {str(m) for m in ws.merged_cells.ranges}
    assert f"A{new_fs}:J{new_fs}" in merges
    assert f"A{new_fs+3}:N{new_fs+3}" in merges
    # summary 行数正确
    assert summary["rows"] == n


# ==================== 回归 ====================
def test_regression_sum_end_tracks_data_not_template(footer_template, blank_pdf, tmp_path):
    """核心回归：合計 SUM 末行跟随真实数据末行、随 N 变化，绝不锁死在模板固定行。"""
    for n in (3, 12):
        _, ws, _ = _run(footer_template, _mk_products(n), blank_pdf, tmp_path)
        data_end, _, new_fs = _layout(n)
        assert ws[f"K{new_fs}"].value == f"=SUM(K{START_ROW}:K{data_end})"
        assert f":K{LAST_ROW})" not in (ws[f"K{new_fs}"].value or ""), "SUM 末行不应还是模板固定的 25"


def test_regression_footer_present_exactly_once(footer_template, blank_pdf, tmp_path):
    """合計行有且仅有一处（不被数据覆盖、也不残留在旧位置）。"""
    n = 4
    _, ws, _ = _run(footer_template, _mk_products(n), blank_pdf, tmp_path)
    _, _, new_fs = _layout(n)
    sum_rows = [r for r in range(1, 60)
                if isinstance(ws.cell(r, 11).value, str) and str(ws.cell(r, 11).value).startswith("=SUM(K")]
    assert sum_rows == [new_fs], f"合計 SUM 应只在第 {new_fs} 行，实际在 {sum_rows}"


# ==================== 边界 / 异常 ====================
def test_boundary_products_exceed_reserved(footer_template, blank_pdf, tmp_path):
    """产品远多于预留段（40 > 8）：页脚自动下移，无覆盖、无崩溃，合計仅一处。"""
    n = 40
    _, ws, _ = _run(footer_template, _mk_products(n), blank_pdf, tmp_path)
    data_end, blank_row, new_fs = _layout(n)
    assert data_end == START_ROW + n - 1
    assert ws[f"B{data_end}"].value  # 第 40 行数据在
    assert all(ws.cell(blank_row, c).value in (None, "") for c in range(1, 15))
    assert ws[f"K{new_fs}"].value == f"=SUM(K{START_ROW}:K{data_end})"
    assert ws[f"A{new_fs}"].value == FOOTER_LABELS[0]


def test_boundary_exactly_fills_reserved(footer_template, blank_pdf, tmp_path):
    """数据正好填满预留段（N=8, data_end==last_row）：页脚下移 1 行。"""
    n = LAST_ROW - START_ROW + 1  # = 8
    _, ws, _ = _run(footer_template, _mk_products(n), blank_pdf, tmp_path)
    data_end, blank_row, new_fs = _layout(n)
    assert data_end == LAST_ROW        # 25
    assert new_fs == FS + 1            # blank=26, 合計=27
    assert ws[f"K{new_fs}"].value == f"=SUM(K{START_ROW}:K{data_end})"


def test_boundary_single_product(footer_template, blank_pdf, tmp_path):
    n = 1
    _, ws, _ = _run(footer_template, _mk_products(n), blank_pdf, tmp_path)
    data_end, blank_row, new_fs = _layout(n)  # data_end=18, blank=19, 合計=20
    assert new_fs == START_ROW + 2
    assert ws[f"K{new_fs}"].value == f"=SUM(K{START_ROW}:K{data_end})"


def test_exception_empty_products_still_raises(footer_template, blank_pdf, tmp_path):
    """空产品仍走 audit 闸门抛 QuoteAuditError（自适应逻辑不吞掉既有异常）。"""
    from skills.drawing_to_quotation import QuoteAuditError
    jp = tmp_path / "empty.json"
    jp.write_text('{"project":"","products":[]}', encoding="utf-8")
    with pytest.raises(QuoteAuditError):
        fill_quote_build(blank_pdf, footer_template, str(jp),
                         str(tmp_path / "e.xlsx"), "", START_ROW, LAST_ROW, False)


# ==================== 端到端（soffice 真重算）====================
def _soffice_works() -> bool:
    if shutil.which("soffice") is None:
        return False
    try:
        with tempfile.TemporaryDirectory() as d:
            src = os.path.join(d, "probe.csv")
            Path(src).write_text("a,b\n1,2\n", encoding="utf-8")
            subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", d, src],
                           check=True, timeout=60,
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            return os.path.exists(os.path.join(d, "probe.pdf"))
    except Exception:
        return False


@pytest.mark.skipif(not _soffice_works(), reason="LibreOffice 不可用")
def test_e2e_soffice_renders_relocated_footer(footer_template, blank_pdf, tmp_path):
    """真让 LibreOffice 打开成品并渲染成 PDF：证明搬迁后的工作簿（含重建的合并/公式）
    有效可渲染，且页脚标签在成品里确实出现（没在搬迁中丢失/损坏）。
    注：LibreOffice headless 默认不重算外部公式，故不校验合計数值，只校验结构可渲染。"""
    n = 6
    out, _, _ = _run(footer_template, _mk_products(n), blank_pdf, tmp_path)
    d = tmp_path / "render"
    d.mkdir()
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(d), out],
                   check=True, timeout=90, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pdf_out = d / (Path(out).stem + ".pdf")
    assert pdf_out.exists(), "LibreOffice 应能把成品渲染成 PDF（工作簿有效、页脚合并未损坏）"
    text = "".join(pg.get_text() for pg in fitz.open(str(pdf_out)))
    for label in (FOOTER_LABELS[0], FOOTER_LABELS[2], "備考"):
        assert label in text, f"渲染成品应含页脚标签「{label}」（搬迁后未丢失）"


# ==================== 黄金快照 ====================
_GOLDEN = Path(__file__).parent / "golden" / "fill_quote_adaptive.json"
_UPDATE = bool(os.environ.get("UPDATE_GOLDEN"))


def _signature(footer_template: str, blank_pdf: str, tmp_path) -> dict:
    """对固定输入（N=6）产出确定性结构签名：布局 + 页脚公式/样式/合并 + 顶部引用。"""
    n = 6
    _, ws, _ = _run(footer_template, _mk_products(n), blank_pdf, tmp_path)
    data_end, blank_row, new_fs = _layout(n)
    footer = []
    for rel in range(4):
        r = new_fs + rel
        footer.append({
            "rel": rel,
            "A": ws[f"A{r}"].value,
            "K": ws[f"K{r}"].value,
            "height": ws.row_dimensions[r].height,
            "bold": bool(ws[f"K{r}"].font.bold),
            "num_fmt": ws[f"K{r}"].number_format,
        })
    return {
        "n": n, "start_row": START_ROW, "last_row": LAST_ROW,
        "data_end": data_end, "blank_row": blank_row, "new_footer_start": new_fs,
        "top_C9": ws["C9"].value,
        "footer": footer,
        "footer_merges": sorted(str(m) for m in ws.merged_cells.ranges
                                if m.min_row >= new_fs),
        "blank_row_empty": all(ws.cell(blank_row, c).value in (None, "") for c in range(1, 15)),
        "old_footer_cleared": ws[f"K{FS}"].value in (None, ""),
    }


def test_golden_fill_quote_adaptive(footer_template, blank_pdf, tmp_path):
    sig = _signature(footer_template, blank_pdf, tmp_path)
    if _UPDATE:
        _GOLDEN.parent.mkdir(exist_ok=True)
        _GOLDEN.write_text(json.dumps(sig, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        pytest.skip("已更新 fill_quote_adaptive.json")
    assert _GOLDEN.exists(), "缺 golden/fill_quote_adaptive.json —— 先跑 UPDATE_GOLDEN=1 生成"
    assert sig == json.loads(_GOLDEN.read_text(encoding="utf-8"))


# ==================== 任务2：数据行显式纯白底 ====================
def _fill_rgb(cell):
    """solid 填充 → ARGB 字符串；无填充 → None。"""
    if cell.fill and cell.fill.patternType:
        rgb = cell.fill.fgColor.rgb
        return rgb if isinstance(rgb, str) else None
    return None


def test_whitefill_data_rows_solid_white(footer_template, blank_pdf, tmp_path):
    """所有数据行 A–N 显式纯白；数量列 I 无填充（去底色）。"""
    _, ws, _ = _run(footer_template, _mk_products(3), blank_pdf, tmp_path)
    for i in range(3):
        r = START_ROW + i
        for col in ("A", "C", "E", "F", "K", "N"):
            assert _fill_rgb(ws[f"{col}{r}"]) == "FFFFFFFF", f"{col}{r} 应为纯白"
        assert _fill_rgb(ws[f"I{r}"]) is None, f"I{r} 数量列应无填充"


def test_whitefill_yellow_overrides_on_flagged_row(footer_template, blank_pdf, tmp_path):
    """⚠行 B/N 淡黄覆盖白底，其余列仍白；I 列仍无填充；未标记行全白。"""
    prods = _mk_products(2)
    prods[1]["note_jp"] = "要確認"          # 触发 ⚠ flag
    _, ws, _ = _run(footer_template, prods, blank_pdf, tmp_path)
    r0, r1 = START_ROW, START_ROW + 1
    # 未标记行：B/N 纯白
    assert _fill_rgb(ws[f"B{r0}"]) == "FFFFFFFF"
    assert _fill_rgb(ws[f"N{r0}"]) == "FFFFFFFF"
    # 标记行：B/N 淡黄；C 仍白；I 仍无填充
    assert _fill_rgb(ws[f"B{r1}"]) == "FFFFF2CC"
    assert _fill_rgb(ws[f"N{r1}"]) == "FFFFF2CC"
    assert _fill_rgb(ws[f"C{r1}"]) == "FFFFFFFF"
    assert _fill_rgb(ws[f"I{r1}"]) is None
